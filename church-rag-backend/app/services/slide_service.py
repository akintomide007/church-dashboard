import os
import io
import json
import textwrap
import unicodedata
import re
from typing import List, Dict, Optional
from pathlib import Path

# Try to import required libraries
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from googleapiclient.discovery import build
    from google.oauth2.service_account import Credentials
    from googleapiclient.http import MediaIoBaseDownload
    GOOGLE_API_AVAILABLE = True
except ImportError:
    GOOGLE_API_AVAILABLE = False

# Configuration
BASE_DIR = Path(__file__).parent.parent.parent
INCOMING_DIR = BASE_DIR / "slide_content" / "incoming"
OUTPUT_DIR = BASE_DIR / "slide_content" / "output_pptx"
SERVICES_FILE = BASE_DIR / "slide_content" / "services.json"
REGISTRY_FILE = BASE_DIR / "slide_content" / "processed_files.json"
CREDS_FILE = BASE_DIR / "slide_content" / "credentials.json"

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

# Default slide generation rules (will be overridden by database preferences)
RULES = {
    "songs": {"lines": 6, "font": 38, "alignment": "center", "font_name": "Arial", "vertical_position": "center"},
    "hymns": {"lines": 8, "font": 34, "alignment": "center", "font_name": "Arial", "vertical_position": "center"},
    "announcements": {"lines": 10, "font": 28, "alignment": "center", "font_name": "Arial", "vertical_position": "center"},
    "uncategorized": {"lines": 8, "font": 32, "alignment": "center", "font_name": "Arial", "vertical_position": "center"},
}

def get_alignment_value(alignment: str):
    """Convert alignment string to python-pptx alignment value (PP_ALIGN constant)."""
    if PPTX_AVAILABLE:
        from pptx.enum.text import PP_ALIGN
        alignment_map = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
            "justify": PP_ALIGN.JUSTIFY
        }
        return alignment_map.get(alignment.lower(), PP_ALIGN.CENTER)
    return 1  # Fallback

def get_vertical_anchor_value(position: str) -> int:
    """Convert vertical position string to python-pptx vertical anchor value."""
    # MSO_ANCHOR values: TOP=1, MIDDLE=3, BOTTOM=4
    position_map = {
        "top": 1,
        "center": 3,
        "middle": 3,
        "bottom": 4
    }
    return position_map.get(position.lower(), 3)  # Default to middle

def load_preferences_from_db():
    """Load slide preferences from database."""
    try:
        from app.core.database import get_db_connection
        conn = get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT songs_lines_per_slide, songs_font_size, songs_text_alignment, songs_font_name, songs_vertical_position,
                   hymns_lines_per_slide, hymns_font_size, hymns_text_alignment, hymns_font_name, hymns_vertical_position,
                   announcements_lines_per_slide, announcements_font_size, announcements_text_alignment, announcements_font_name, announcements_vertical_position,
                   uncategorized_lines_per_slide, uncategorized_font_size, uncategorized_text_alignment, uncategorized_font_name, uncategorized_vertical_position
            FROM slide_preferences
            WHERE user_id = 1
            LIMIT 1
        """)
        
        result = cursor.fetchone()
        cursor.close()
        conn.close()
        
        if result:
            return {
                "songs": {
                    "lines": result[0],
                    "font": result[1],
                    "alignment": result[2] or "center",
                    "font_name": result[3] or "Arial",
                    "vertical_position": result[4] or "center"
                },
                "hymns": {
                    "lines": result[5],
                    "font": result[6],
                    "alignment": result[7] or "center",
                    "font_name": result[8] or "Arial",
                    "vertical_position": result[9] or "center"
                },
                "announcements": {
                    "lines": result[10],
                    "font": result[11],
                    "alignment": result[12] or "center",
                    "font_name": result[13] or "Arial",
                    "vertical_position": result[14] or "center"
                },
                "uncategorized": {
                    "lines": result[15],
                    "font": result[16],
                    "alignment": result[17] or "center",
                    "font_name": result[18] or "Arial",
                    "vertical_position": result[19] or "center"
                }
            }
    except Exception as e:
        print(f"Could not load preferences from database: {e}")
    
    return RULES  # Return defaults if database load fails

# -----------------------------
# Utility functions
# -----------------------------
def safe_filename(name: str) -> str:
    """Sanitize filenames for cross-platform compatibility."""
    # Normalize Unicode to ASCII
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    # Replace unsafe characters
    name = re.sub(r'[<>:"/\\|?*]', "_", name)
    # Replace smart and straight quotes with underscore
    name = re.sub(r"['']", "_", name)
    return name.strip()

def ensure_directories():
    """Ensure all required directories exist."""
    INCOMING_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    SERVICES_FILE.parent.mkdir(parents=True, exist_ok=True)

# -----------------------------
# Services management
# -----------------------------
def load_services() -> List[Dict]:
    """Load services configuration from JSON file."""
    ensure_directories()
    
    if not SERVICES_FILE.exists():
        # Create default services file
        default_services = {
            "services": [
                {
                    "name": "08AM",
                    "start_time": "08:00",
                    "drive_folder_id": ""
                },
                {
                    "name": "11AM",
                    "start_time": "11:00",
                    "drive_folder_id": ""
                },
                {
                    "name": "06PM",
                    "start_time": "18:00",
                    "drive_folder_id": ""
                }
            ]
        }
        with open(SERVICES_FILE, "w") as f:
            json.dump(default_services, f, indent=2)
        return default_services["services"]
    
    with open(SERVICES_FILE, "r") as f:
        return json.load(f)["services"]

def get_service(name: str) -> Optional[Dict]:
    """Get a specific service by name."""
    for service in load_services():
        if service["name"] == name:
            return service
    return None

# -----------------------------
# Registry management
# -----------------------------
def load_registry() -> Dict:
    """Load the registry of processed files."""
    if not REGISTRY_FILE.exists():
        return {"processed_ids": []}
    with open(REGISTRY_FILE, "r") as f:
        return json.load(f)

def save_registry(data: Dict):
    """Save the registry of processed files."""
    ensure_directories()
    with open(REGISTRY_FILE, "w") as f:
        json.dump(data, f, indent=2)

# -----------------------------
# Google Drive functions
# -----------------------------
def get_drive_service():
    """Create and return Google Drive API service."""
    if not GOOGLE_API_AVAILABLE:
        raise ImportError("Google API libraries not installed. Install: pip install google-api-python-client google-auth")
    
    if not CREDS_FILE.exists():
        raise FileNotFoundError(
            f"credentials.json not found at {CREDS_FILE}. "
            "Please create a Google Service Account and place credentials.json in the slide_content directory."
        )
    
    creds = Credentials.from_service_account_file(str(CREDS_FILE), scopes=SCOPES)
    return build("drive", "v3", credentials=creds)

def fetch_and_download(service: Dict) -> str:
    """
    Download files from Google Drive for a service.
    Returns a status message.
    """
    ensure_directories()
    
    service_name = service["name"]
    folder_id = service.get("drive_folder_id", "")
    
    if not folder_id:
        return f"No Google Drive folder configured for service {service_name}"
    
    registry = load_registry()
    drive = get_drive_service()
    
    query = f"'{folder_id}' in parents and trashed=false"
    results = drive.files().list(
        q=query,
        fields="files(id, name, mimeType, parents, createdTime)"
    ).execute()
    
    files = results.get("files", [])
    if not files:
        return f"No new files found for {service_name}."
    
    downloaded_count = 0
    category = "uncategorized"
    target_dir = INCOMING_DIR / service_name / category
    target_dir.mkdir(parents=True, exist_ok=True)
    
    for file in files:
        if file["id"] in registry["processed_ids"]:
            continue
        
        safe_name = safe_filename(file["name"])
        file_path = target_dir / safe_name
        
        try:
            request = drive.files().get_media(fileId=file["id"])
            fh = io.FileIO(str(file_path), "wb")
            downloader = MediaIoBaseDownload(fh, request)
            
            done = False
            while not done:
                _, done = downloader.next_chunk()
            
            registry["processed_ids"].append(file["id"])
            downloaded_count += 1
        except Exception as e:
            print(f"Error downloading {file['name']}: {str(e)}")
            continue
    
    save_registry(registry)
    
    if downloaded_count > 0:
        return f"Successfully fetched {downloaded_count} file(s) for service {service_name}"
    else:
        return f"No new files to download for {service_name}"

# -----------------------------
# Slide generation functions
# -----------------------------
def read_file(path: Path) -> str:
    """Read content from a text or Word document."""
    if path.suffix == ".txt":
        with open(path, encoding="utf-8") as f:
            return f.read()
    
    if path.suffix == ".docx":
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx not installed. Install: pip install python-docx")
        
        try:
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            # Fallback to plain text if Word doc is corrupted
            try:
                with open(path, encoding="utf-8") as f:
                    return f.read()
            except Exception as e:
                print(f"Error reading {path.name}: {e}")
                return ""
    
    return ""

def chunk_text(lines: List[str], max_lines: int, max_chars: int = 35) -> List[List[str]]:
    """Chunk text into slides based on line limits."""
    slides, current = [], []
    for line in lines:
        for wrapped in textwrap.wrap(line, max_chars):
            current.append(wrapped)
            if len(current) == max_lines:
                slides.append(current)
                current = []
    if current:
        slides.append(current)
    return slides

def generate_slides(service_name: str) -> str:
    """
    Generate PowerPoint slides for a service.
    Returns a status message.
    """
    if not PPTX_AVAILABLE:
        raise ImportError("python-pptx not installed. Install: pip install python-pptx")
    
    ensure_directories()
    
    base_path = INCOMING_DIR / service_name
    if not base_path.exists():
        return f"No content found for service {service_name}"
    
    # Load preferences from database
    rules = load_preferences_from_db()
    
    generated_count = 0
    categories = set([d.name for d in base_path.iterdir() if d.is_dir()])
    categories.add("uncategorized")
    
    for category in categories:
        cat_path = base_path / category
        if not cat_path.exists():
            continue
        
        rule = rules.get(category, rules["uncategorized"])
        
        for file_path in cat_path.iterdir():
            if not file_path.is_file():
                continue
            
            # Create a separate presentation for each file
            prs = Presentation()
            
            text = read_file(file_path)
            if not text.strip():
                continue
            
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            slides = chunk_text(lines, rule["lines"])
            
            for slide_lines in slides:
                # Use blank layout (index 6)
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                
                # Add textbox - full slide dimensions for flexibility
                # Standard slide dimensions: 10" x 7.5"
                left = Inches(0.5)
                top = Inches(0.5)
                width = Inches(9)
                height = Inches(6.5)
                
                box = slide.shapes.add_textbox(left, top, width, height)
                tf = box.text_frame
                tf.word_wrap = True
                
                # Apply vertical positioning
                vertical_pos = rule.get("vertical_position", "center")
                tf.vertical_anchor = get_vertical_anchor_value(vertical_pos)
                tf.clear()
                
                # Add content
                for i, line in enumerate(slide_lines):
                    if i == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    
                    p.text = line
                    p.font.size = Pt(rule["font"])
                    p.font.name = rule.get("font_name", "Arial")
                    p.font.bold = False  # Not bold
                    p.alignment = get_alignment_value(rule.get("alignment", "center"))
            
            # Save with filename
            base_name = file_path.stem
            out_filename = f"{service_name}_{category}_{base_name}.pptx"
            output_path = OUTPUT_DIR / out_filename
            prs.save(str(output_path))
            generated_count += 1
    
    return f"Generated {generated_count} PowerPoint file(s) for {service_name}"

def list_generated_slides(service_filter: Optional[str] = None) -> List[Dict]:
    """List all generated PowerPoint files."""
    ensure_directories()
    
    if not OUTPUT_DIR.exists():
        return []
    
    files = []
    for file_path in OUTPUT_DIR.glob("*.pptx"):
        # Parse filename: SERVICE_CATEGORY_NAME.pptx
        parts = file_path.stem.split("_", 2)
        if len(parts) >= 3:
            service, category, name = parts[0], parts[1], parts[2]
        elif len(parts) == 2:
            service, category, name = parts[0], "uncategorized", parts[1]
        else:
            service, category, name = "unknown", "uncategorized", parts[0]
        
        if service_filter and service != service_filter:
            continue
        
        files.append({
            "filename": file_path.name,
            "service": service,
            "category": category,
            "path": str(file_path.relative_to(BASE_DIR))
        })
    
    return sorted(files, key=lambda x: (x["service"], x["category"], x["filename"]))
